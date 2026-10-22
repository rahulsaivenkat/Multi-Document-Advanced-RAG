"""Shared pytest fixtures, synthetic document generators, and test client setup."""
import io
import os
import shutil
import sys
from pathlib import Path

# Ensure project root is in sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from typing import Generator
import pytest
from fastapi.testclient import TestClient

# Configure test environment variables before importing app modules
os.environ["CHROMA_PERSIST_DIRECTORY"] = "data/test_storage/chroma"
os.environ["BM25_PERSIST_DIRECTORY"] = "data/test_storage/bm25"
os.environ["UPLOAD_DIRECTORY"] = "data/test_uploads"
os.environ["GROQ_API_KEY"] = "mock_groq_api_key_for_testing"

from app.config import settings
from app.main import app
from app.schemas import Chunk, SourceLocation
from app.services.ingestion import IngestionService
from app.services.retrieval import RetrievalPipeline


@pytest.fixture(scope="session", autouse=True)
def setup_test_directories():
    """Create test directories and clean them up after test session."""
    test_dirs = ["data/test_storage", "data/test_uploads"]
    for d in test_dirs:
        Path(d).mkdir(parents=True, exist_ok=True)
    yield
    for d in test_dirs:
        if Path(d).exists():
            shutil.rmtree(d, ignore_errors=True)


@pytest.fixture
def client() -> Generator[TestClient, None, None]:
    """FastAPI TestClient fixture."""
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    """Generate a multi-page PDF fixture with extractable text."""
    import pymupdf as fitz
    pdf_path = tmp_path / "test_report.pdf"
    doc = fitz.open()

    # Page 1
    page1 = doc.new_page()
    page1.insert_text((50, 72), "PaperMind AI Specification Overview.\nThis is the executive summary on Page 1.")

    # Page 2
    page2 = doc.new_page()
    page2.insert_text((50, 72), "Technical Details on Page 2.\nChromaDB and BM25 form the hybrid search engine.")

    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def sample_scanned_pdf(tmp_path: Path) -> Path:
    """Generate an image-only / textless PDF to verify OCR limitation error."""
    import pymupdf as fitz
    pdf_path = tmp_path / "scanned_doc.pdf"
    doc = fitz.open()
    doc.new_page()  # Blank page without text
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def sample_txt(tmp_path: Path) -> Path:
    """Generate a plain text fixture with multiple paragraphs and lines."""
    txt_path = tmp_path / "notes.txt"
    content = (
        "Paragraph One: Architecture design principles for RAG.\n"
        "Ensure verifiable citations and strict evidence grounding.\n\n"
        "Paragraph Two: Evaluation metrics.\n"
        "Measure Recall@k and Mean Reciprocal Rank."
    )
    txt_path.write_text(content, encoding="utf-8")
    return txt_path


@pytest.fixture
def sample_md(tmp_path: Path) -> Path:
    """Generate a Markdown fixture with heading hierarchy and code block."""
    md_path = tmp_path / "guide.md"
    content = (
        "# User Guide\n\n"
        "Introduction to PaperMind AI.\n\n"
        "## Installation\n\n"
        "### Windows Setup\n\n"
        "Run the following command:\n\n"
        "```powershell\n"
        "python -m venv .venv\n"
        "pip install -r requirements.txt\n"
        "```\n\n"
        "Ensure Python 3.10+ is installed."
    )
    md_path.write_text(content, encoding="utf-8")
    return md_path


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    """Generate a DOCX fixture with paragraphs, headings, and a table."""
    import docx
    docx_path = tmp_path / "document.docx"
    doc = docx.Document()
    doc.add_heading("System Architecture", level=1)
    doc.add_paragraph("The hybrid search engine merges dense vectors and BM25 scores.")
    
    # Add a table
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Component"
    table.rows[0].cells[1].text = "Technology"
    table.rows[1].cells[0].text = "Vector Store"
    table.rows[1].cells[1].text = "ChromaDB"

    doc.save(str(docx_path))
    return docx_path


@pytest.fixture
def sample_csv(tmp_path: Path) -> Path:
    """Generate a CSV fixture with headers and rows."""
    csv_path = tmp_path / "employees.csv"
    content = (
        "id,name,role,department\n"
        "101,Alice Smith,Principal Engineer,AI Research\n"
        "102,Bob Johnson,Systems Architect,Platform\n"
        "103,Carol White,Security Analyst,Infosec\n"
    )
    csv_path.write_text(content, encoding="utf-8")
    return csv_path


@pytest.fixture
def sample_html(tmp_path: Path) -> Path:
    """Generate an HTML fixture with headings, paragraphs, table, and noise tags."""
    html_path = tmp_path / "manual.html"
    content = (
        "<!DOCTYPE html><html><head><title>Test Manual</title>"
        "<script>console.log('malicious script');</script>"
        "<style>body { color: red; }</style>"
        "</head><body>"
        "<nav><a href='#'>Home</a></nav>"
        "<h1>PaperMind Manual</h1>"
        "<p>This manual explains the retrieval pipeline.</p>"
        "<h2>Components</h2>"
        "<ul><li>SentenceTransformers</li><li>CrossEncoder</li></ul>"
        "<table><tr><th>Module</th><th>Status</th></tr><tr><td>RRF</td><td>Active</td></tr></table>"
        "<footer>Copyright 2026</footer>"
        "</body></html>"
    )
    html_path.write_text(content, encoding="utf-8")
    return html_path


@pytest.fixture
def sample_json(tmp_path: Path) -> Path:
    """Generate a JSON fixture with nested objects and arrays."""
    json_path = tmp_path / "config.json"
    content = (
        '{\n'
        '  "cluster": "prod-east",\n'
        '  "settings": {\n'
        '    "timeout": 30,\n'
        '    "retry": true\n'
        '  },\n'
        '  "servers": [\n'
        '    {"id": "s1", "host": "10.0.0.1", "role": "primary"},\n'
        '    {"id": "s2", "host": "10.0.0.2", "role": "replica"}\n'
        '  ]\n'
        '}'
    )
    json_path.write_text(content, encoding="utf-8")
    return json_path


@pytest.fixture
def sample_python(tmp_path: Path) -> Path:
    """Generate a Python source code fixture with class, methods, and functions."""
    py_path = tmp_path / "calculator.py"
    content = (
        '"""Module providing math utilities."""\n\n'
        'import math\n\n'
        'class Calculator:\n'
        '    """Arithmetic calculator class."""\n\n'
        '    def add(self, a: float, b: float) -> float:\n'
        '        """Add two numbers."""\n'
        '        return a + b\n\n'
        '    def multiply(self, a: float, b: float) -> float:\n'
        '        """Multiply two numbers."""\n'
        '        return a * b\n\n'
        'def compute_square_root(x: float) -> float:\n'
        '    """Compute square root of a non-negative number."""\n'
        '    return math.sqrt(x)\n'
    )
    py_path.write_text(content, encoding="utf-8")
    return py_path
