"""Word (.docx) document parser extracting paragraphs, headings, and tables with stable coordinates."""
from pathlib import Path
from typing import List, Optional
import docx
from app.parsers.base import BaseParser, ParserError
from app.schemas import ExtractedBlock, SourceLocation


class DocxParser(BaseParser):
    """Parses DOCX files into paragraph and table blocks without fabricating page numbers."""

    def parse(self, file_path: Path, document_id: str, source_filename: str) -> List[ExtractedBlock]:
        self.check_non_empty(file_path)

        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise ParserError(f"Failed to open or parse DOCX file '{source_filename}': {str(e)}")

        blocks: List[ExtractedBlock] = []
        heading_stack: List[str] = []
        block_counter = 1

        def current_heading_path() -> Optional[str]:
            return " > ".join(heading_stack) if heading_stack else None

        # 1. Extract Paragraphs in order
        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            style_name = getattr(p.style, "name", "") if p.style else ""
            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.replace("Heading", "").strip())
                except ValueError:
                    level = 1
                
                # Adjust stack
                while len(heading_stack) >= level:
                    heading_stack.pop()
                heading_stack.append(text)

            blocks.append(
                ExtractedBlock(
                    document_id=document_id,
                    source_filename=source_filename,
                    file_type=".docx",
                    content=text,
                    block_id=f"para_{p_idx}",
                    source_location=SourceLocation(
                        paragraph_index=p_idx,
                        heading_path=current_heading_path()
                    ),
                    metadata={"style": style_name}
                )
            )
            block_counter += 1

        # 2. Extract Tables
        for t_idx, table in enumerate(doc.tables):
            # Read header row if available
            headers = []
            if len(table.rows) > 0:
                headers = [cell.text.strip() for cell in table.rows[0].cells]

            for r_idx, row in enumerate(table.rows):
                row_cells = [cell.text.strip() for cell in row.cells]
                # Skip completely blank rows
                if not any(row_cells):
                    continue

                # Format row text nicely with headers if available
                if headers and r_idx > 0:
                    formatted_pairs = [f"{h}: {c}" if h else c for h, c in zip(headers, row_cells)]
                    row_content = " | ".join(formatted_pairs)
                else:
                    row_content = " | ".join(row_cells)

                blocks.append(
                    ExtractedBlock(
                        document_id=document_id,
                        source_filename=source_filename,
                        file_type=".docx",
                        content=f"[Table {t_idx + 1}, Row {r_idx + 1}] {row_content}",
                        block_id=f"table_{t_idx}_row_{r_idx}",
                        source_location=SourceLocation(
                            table_index=t_idx,
                            row_number=r_idx + 1,
                            column_names=headers if headers else None,
                            heading_path=current_heading_path()
                        ),
                        metadata={"is_table": True, "table_index": t_idx, "row_index": r_idx}
                    )
                )
                block_counter += 1

        if not blocks:
            raise ParserError(f"DOCX document '{source_filename}' contains no readable paragraphs or tables.")

        return blocks
